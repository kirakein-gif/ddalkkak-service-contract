"""통학차량 계약 문서 생성기.

Rule Engine의 판정 결과와 학교 입력값을 조합하여 다음 문서를 생성한다.

1. 견적제출 안내공고 / 입찰공고 / 1인견적 검토서
2. 과업지시서
3. 계약 특수조건
4. 운행노선표

문구는 학교 실무에서 반복되는 공통 구조를 재구성한 초안이며,
학교별 내부규정·운영계획·공고일 기준 최신 법령을 최종 확인해야 한다.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

from app.rules.transport import TransportRuleResult


@dataclass(slots=True)
class RouteEntry:
    route_name: str
    direction: str = "등교"
    departure: str = ""
    via: str = ""
    destination: str = ""
    departure_time: str = ""
    notes: str = ""


@dataclass(slots=True)
class TransportDocumentData:
    school_name: str
    service_name: str
    estimated_price: int
    vehicle_count: int
    seat_capacity_min: int
    planned_date: date
    base_amount: int | None = None
    budget_amount: int | None = None
    contract_start: date | None = None
    contract_end: date | None = None
    operation_days: int | None = None
    school_address: str = ""
    contact_name: str = ""
    contact_phone: str = ""
    notice_number: str = ""
    region_restriction_text: str = ""
    bid_start: str = ""
    bid_end: str = ""
    bid_open: str = ""
    vehicle_year_condition: str = ""
    direct_vehicle_required: bool = False
    pricing_method_label: str = "총액"
    service_item_name: str = ""
    service_item_code: str | None = None
    driver_included: bool = True
    attendant_included: bool = False
    routes: list[RouteEntry] = field(default_factory=list)
    extra_notes: str = ""


@dataclass(slots=True)
class GeneratedDocument:
    key: str
    title: str
    content: str


def _money(value: int | None) -> str:
    if value is None:
        return "[입력 필요]"
    return f"{value:,}원"


def _date_range(start: date | None, end: date | None) -> str:
    if not start or not end:
        return "[계약기간 입력 필요]"
    return f"{start.isoformat()} ~ {end.isoformat()}"


def _value(value: str, fallback: str = "[입력 필요]") -> str:
    return value.strip() if value and value.strip() else fallback


def _notice_title(rule: TransportRuleResult) -> str:
    method = rule.contract_method
    if "1인 견적" in method:
        return "통학차량 임차용역 수의계약 검토 및 견적요청서"
    if "2인 이상" in method or "수의계약" in method:
        return "통학차량 임차용역 소액수의 견적제출 안내공고"
    return "통학차량 임차용역 입찰공고"


def _qualification_lines(data: TransportDocumentData, rule: TransportRuleResult) -> list[str]:
    lines = [
        "지방계약 관계 법령 및 공고조건에 따라 입찰·견적 제출 자격을 갖춘 업체",
        "여객자동차운송사업 중 전세버스운송사업 등록 업체",
    ]
    if rule.service_item_code:
        lines.append(
            f"나라장터 세부품명 {rule.service_item_name}({rule.service_item_code}) 관련 자격요건 충족 업체"
        )
    if "필요" in rule.direct_production_status:
        lines.append(
            f"{rule.service_item_name}({rule.service_item_code}) 직접생산확인증명서를 유효하게 보유한 업체"
        )
    if data.region_restriction_text:
        lines.append(f"지역제한: {data.region_restriction_text}")
    return lines


def _build_notice(data: TransportDocumentData, rule: TransportRuleResult) -> GeneratedDocument:
    title = _notice_title(rule)
    qualifications = "\n".join(f"- {line}" for line in _qualification_lines(data, rule))
    rate = (
        f"{rule.minimum_quote_rate * 100:.3f}% 이상"
        if rule.minimum_quote_rate is not None
        else "공고일 기준 적용기준 확인"
    )
    qualification = (
        f"{rule.qualification_score}점 이상"
        if rule.qualification_score is not None
        else "해당 없음 또는 별도 확인"
    )

    content = f"""# {title}

{data.school_name}

## 1. 견적·입찰에 부치는 사항
- 용역명: {data.service_name}
- 용역기간: {_date_range(data.contract_start, data.contract_end)}
- 차량규격: 최소 {data.seat_capacity_min}인승, {data.vehicle_count}대
- 예정 운행일수: {data.operation_days if data.operation_days else '[입력 필요]'}일
- 계약방식: {data.pricing_method_label}계약
- 추정가격: {_money(data.estimated_price)}
- 기초금액: {_money(data.base_amount)}
- 예산액: {_money(data.budget_amount)}
- 세부품명: {_value(data.service_item_name)}{f' ({data.service_item_code})' if data.service_item_code else ''}

## 2. 계약 및 낙찰방법
- 계약경로: {rule.contract_method}
- 전자조달: {'원칙적 사용' if rule.designated_system_required else '필수 아님'}
- 낙찰·견적가격 기준: {rate}
- 적격심사 통과기준: {qualification}
- 지역제한 검토: {rule.regional_restriction_status}
- 직접생산 확인: {rule.direct_production_status}

## 3. 참가자격
{qualifications}

## 4. 견적·입찰 일정
- 제출개시: {_value(data.bid_start)}
- 제출마감: {_value(data.bid_end)}
- 개찰일시: {_value(data.bid_open)}
- 개찰장소: 국가종합전자조달시스템 또는 공고에서 정한 장소

## 5. 계약상대자 결정
- 예정가격 및 계약상대자 결정은 공고일 현재 적용되는 지방계약 관계 규정과 행정안전부 예규를 따른다.
- 적격심사 대상인 경우 공고일 현재 유효한 조달청 일반용역 적격심사 세부기준을 적용한다.
- 동일가격 제출자가 발생하는 경우 전자조달시스템의 자동추첨 등 관계 규정에 따른다.
- 수의계약 대상인 경우 수의계약 배제사유와 이해충돌방지법상 체결 제한을 확인한다.

## 6. 계약 전 확인서류
{chr(10).join(f'- {item}' for item in rule.required_documents)}

## 7. 유의사항
- 과업지시서, 계약 특수조건 및 운행노선표를 반드시 함께 확인한다.
- 학교가 제시하는 차량연식·직영차량 등 선택조건은 공고 전 경쟁제한의 적정성을 확인한다.
- 계약상대자는 관계 법령상 안전기준과 학교의 안전관리 요구사항을 준수해야 한다.
- 공고일 현재 법령·예규·조달청 기준·SMPP 중소기업자간 경쟁제품 정보를 최종 확인한다.

## 8. 문의
- 기관: {data.school_name}
- 주소: {_value(data.school_address)}
- 담당자: {_value(data.contact_name)}
- 연락처: {_value(data.contact_phone)}
- 공고번호: {_value(data.notice_number)}

※ 본 문서는 딸깍 용역계약이 생성한 검토용 초안입니다. 최종 공고 전 담당자가 원문 규정과 기관 지침을 확인해야 합니다.
"""
    return GeneratedDocument("notice", title, content)


def _build_scope(data: TransportDocumentData, rule: TransportRuleResult) -> GeneratedDocument:
    title = f"{data.service_name} 과업지시서"
    direct = "원칙적으로 계약상대자 소유·직영차량을 사용" if data.direct_vehicle_required else "차량 소유·사용관계를 확인하고 공고조건에 적합한 차량을 사용"
    year_condition = _value(data.vehicle_year_condition, "법정 차령을 준수하고 학교가 별도로 정한 연식조건이 있는 경우 이를 충족")

    content = f"""# {title}

## 1. 과업 목적
본 과업은 {data.school_name} 학생의 안전하고 원활한 통학을 위해 통학차량과 필요한 운송서비스를 제공하는 것을 목적으로 한다.

## 2. 과업 개요
- 용역명: {data.service_name}
- 계약기간: {_date_range(data.contract_start, data.contract_end)}
- 차량대수: {data.vehicle_count}대
- 차량정원: 차량별 최소 {data.seat_capacity_min}인승
- 예정 운행일수: {data.operation_days if data.operation_days else '[입력 필요]'}일
- 운전원 포함: {'예' if data.driver_included else '아니오'}
- 동승보호자 포함: {'예' if data.attendant_included else '아니오'}
- 세부품명: {_value(data.service_item_name)}{f' ({data.service_item_code})' if data.service_item_code else ''}

## 3. 차량 조건
- 전세버스운송사업 등록 등 관계 법령상 운송자격을 갖춘 차량을 운행한다.
- {direct}.
- 차량연식: {year_condition}.
- 차량은 관계 법령상 차령, 정기검사, 보험 및 안전기준을 충족해야 한다.
- 좌석안전띠, 냉·난방, 후방확인장치 등 운행에 필요한 안전·편의설비를 정상 상태로 유지한다.
- 어린이통학버스에 해당하는 차량은 신고, 표시, 구조 및 안전장치 기준을 충족해야 한다.

## 4. 운전원 및 동승보호자
- 운전원은 적법한 운전면허와 버스운전자격을 갖추고 관계 법령상 결격사유가 없어야 한다.
- 운전원 교체 시 학교에 사전 통보하고 필요한 증빙을 제출한다.
- 어린이통학버스 동승보호자가 필요한 경우 성년 보호자를 배치하고 안전교육 이수 여부를 확인한다.
- 학교의 학생안전 관련 내부절차에 따라 필요한 범죄경력·아동학대 관련 확인 절차에 협조한다.

## 5. 운행
- 학교가 정한 운행노선, 운행시간 및 승·하차 지점을 준수한다.
- 운행개시 전 노선확인 또는 시험운행이 필요한 경우 학교와 협의하여 실시한다.
- 교통상황, 학사일정 또는 학생 수 변동으로 노선·시간 변경이 필요한 경우 학교의 승인을 받아 조정한다.
- 임의 결행, 임의 노선변경 또는 학교 승인 없는 차량교체를 하지 않는다.

## 6. 안전관리
- 운행 전 차량의 제동장치, 타이어, 등화장치, 안전띠 등 주요 안전상태를 점검한다.
- 사고·고장 발생 시 즉시 학교에 알리고 학생안전을 우선하여 대체차량 등 필요한 조치를 한다.
- 어린이통학버스 해당 시 하차확인장치 등 법정 안전장치를 정상 작동시킨다.
- 운전자와 동승보호자는 학생 승·하차 시 안전을 확인한다.

## 7. 비용부담
계약금액에는 별도 명시가 없는 한 차량운행에 필요한 유류비, 정비비, 보험료, 인건비, 일반관리비 및 이윤 등 통상적인 운행 제경비를 포함하는 것으로 한다. 통행료·주차료 등 별도 정산항목은 공고·산출내역서에서 구분한다.

## 8. 제출 및 관리서류
{chr(10).join(f'- {item}' for item in rule.required_documents)}

## 9. 기타
- 본 과업지시서와 공고문·특수조건이 상충하는 경우 계약담당자가 관계 규정과 공고취지에 따라 해석·정리한다.
- 학교의 학사일정 변경 등 불가피한 사유가 발생하면 계약당사자 간 협의하여 관계 규정에 따라 변경한다.
- 추가사항: {_value(data.extra_notes, '없음')}

※ 본 문서는 검토용 초안이며 학교별 운영계획과 최신 법령을 반영하여 확정해야 합니다.
"""
    return GeneratedDocument("scope", title, content)


def _build_special_conditions(data: TransportDocumentData, rule: TransportRuleResult) -> GeneratedDocument:
    title = f"{data.service_name} 계약 특수조건"
    conditions = "\n".join(f"- {item}" for item in rule.common_school_conditions)
    content = f"""# {title}

## 제1조 목적
이 특수조건은 {data.school_name}과 계약상대자 사이의 {data.service_name} 계약을 안전하고 성실하게 이행하기 위해 필요한 사항을 정한다.

## 제2조 계약문서
공고문, 과업지시서, 운행노선표, 산출내역서 및 이 특수조건은 계약문서의 일부로 본다.

## 제3조 차량 및 운행
- 계약상대자는 계약조건에 적합한 차량 {data.vehicle_count}대를 확보하여 운행한다.
- 학교의 사전승인 없이 차량·운전원을 임의 변경하지 않는다.
- 차량 교체가 필요한 경우 동급 이상의 안전조건을 갖춘 차량을 투입하고 관련 서류를 제출한다.
- 운행노선과 시간은 학교의 학사운영에 맞추어 준수한다.

## 제4조 안전 및 보험
- 관계 법령에 따른 자동차보험 및 학교가 요구한 보험조건을 계약기간 동안 유지한다.
- 차량의 검사·정비 상태를 상시 유지하고 사고 예방에 필요한 조치를 한다.
- 어린이통학버스에 해당하는 경우 신고·안전장치·동승보호자·안전교육 등 법정의무를 이행한다.

## 제5조 운전원 및 동승보호자
- 운전원은 면허·버스운전자격 등 필요한 자격을 유지한다.
- 동승보호자가 계약범위에 포함된 경우 적격한 성년 보호자를 배치한다.
- 인력 교체 시 학교가 요구하는 확인서류를 제출하고 안전업무 인수인계를 실시한다.

## 제6조 서류제출
계약상대자는 공고와 과업지시서에서 요구한 차량·보험·자격·직접생산·안전 관련 서류를 정해진 기한까지 제출한다.

## 제7조 사고·고장 및 대체운행
- 사고 또는 고장 발생 시 즉시 학교에 통보하고 학생 안전을 최우선으로 조치한다.
- 운행이 불가능한 경우 계약조건에 부합하는 대체차량을 신속히 투입한다.
- 사고처리 및 보험절차에 성실히 협조한다.

## 제8조 금지사항
- 학교의 승인 없이 계약상 운송업무를 제3자에게 임의로 전가하지 않는다.
- 음주·과로·난폭운전 등 학생안전을 저해하는 행위를 금한다.
- 운행 중 학생 개인정보나 학교 내부정보를 목적 외로 사용하지 않는다.

## 제9조 학교 선택 특수조건 후보
{conditions}
- 차량연식 조건: {_value(data.vehicle_year_condition, '별도 조건 없음')}
- 직영차량 요구: {'적용' if data.direct_vehicle_required else '미적용 또는 별도 검토'}

## 제10조 계약변경 및 해석
학사일정·학생수·노선 등 계약조건의 변경이 필요한 경우 관계 법령과 계약조건에 따라 상호 협의하여 처리한다.

※ 본 특수조건은 자동생성 초안입니다. 학교가 선택하는 조건은 경쟁제한의 적정성을 검토한 뒤 확정해야 합니다.
"""
    return GeneratedDocument("special_conditions", title, content)


def _build_routes(data: TransportDocumentData) -> GeneratedDocument:
    title = f"{data.service_name} 운행노선표"
    if data.routes:
        rows = []
        for idx, route in enumerate(data.routes, start=1):
            rows.append(
                f"{idx}. [{route.direction}] {route.route_name}\n"
                f"   - 출발: {_value(route.departure)}\n"
                f"   - 경유: {_value(route.via, '없음')}\n"
                f"   - 도착: {_value(route.destination)}\n"
                f"   - 출발시간: {_value(route.departure_time)}\n"
                f"   - 비고: {_value(route.notes, '없음')}"
            )
        route_text = "\n\n".join(rows)
    else:
        route_text = """1. [등교] 1호차
   - 출발: [입력 필요]
   - 경유: [입력 필요]
   - 도착: 학교
   - 출발시간: [입력 필요]
   - 비고: [입력 필요]

2. [하교] 1호차
   - 출발: 학교
   - 경유: [입력 필요]
   - 도착: [입력 필요]
   - 출발시간: [입력 필요]
   - 비고: [입력 필요]"""

    content = f"""# {title}

- 학교명: {data.school_name}
- 용역명: {data.service_name}
- 운행기간: {_date_range(data.contract_start, data.contract_end)}
- 차량대수: {data.vehicle_count}대
- 최소 승차정원: {data.seat_capacity_min}인승
- 예정 운행일수: {data.operation_days if data.operation_days else '[입력 필요]'}일

## 운행노선

{route_text}

## 노선 운영 원칙
- 승·하차 장소와 시간은 학생 안전을 우선하여 정한다.
- 도로상황·학사일정·학생 변동으로 조정이 필요한 경우 학교 승인을 받아 변경한다.
- 계약상대자는 운행개시 전 실제 노선을 확인하고 차량 통행 가능 여부와 소요시간을 점검한다.
- 정원초과 운행을 하지 않는다.

※ 최종 노선은 학교의 학생 배치현황과 학사일정을 반영하여 확정합니다.
"""
    return GeneratedDocument("routes", title, content)


def build_transport_documents(
    data: TransportDocumentData,
    rule: TransportRuleResult,
) -> list[GeneratedDocument]:
    return [
        _build_notice(data, rule),
        _build_scope(data, rule),
        _build_special_conditions(data, rule),
        _build_routes(data),
    ]


def _set_default_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(10.5)


def _write_markdown_like(document: Document, text: str) -> None:
    for raw in text.splitlines():
        line = raw.rstrip()
        if not line:
            document.add_paragraph()
            continue
        if line.startswith("# "):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:])
            run.bold = True
            run.font.size = Pt(16)
            continue
        if line.startswith("## "):
            p = document.add_paragraph()
            run = p.add_run(line[3:])
            run.bold = True
            run.font.size = Pt(12)
            continue
        if line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
            continue
        document.add_paragraph(line)


def document_to_docx_bytes(document_data: GeneratedDocument) -> bytes:
    document = Document()
    _set_default_font(document)
    _write_markdown_like(document, document_data.content)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def build_transport_docx_zip(documents: list[GeneratedDocument]) -> bytes:
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        for index, doc in enumerate(documents, start=1):
            filename = f"{index:02d}_{doc.title.replace('/', '_')}.docx"
            archive.writestr(filename, document_to_docx_bytes(doc))
    return stream.getvalue()
