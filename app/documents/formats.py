"""공통 계약문서 출력 포맷 유틸리티.

공사·용역·물품 모든 모듈이 HWPX/DOCX 패키징에 재사용한다.
공고문은 실제 학교·나라장터 공고문에 가까운 전용 HWPX 레이아웃으로 출력한다.
"""

from dataclasses import dataclass, field
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from hwpx import HwpxDocument


@dataclass(slots=True)
class GeneratedDocument:
    key: str
    title: str
    content: str
    layout: str = "default"
    notice_number: str = ""
    issuer: str = ""
    summary_rows: tuple[tuple[str, str], ...] = field(default_factory=tuple)
    issue_date: str = ""
    signatory: str = ""
    alert_text: str = ""


def _safe_filename(title: str) -> str:
    return title.replace("/", "_").replace("\\", "_").replace(":", "_")


def _paragraph_alignment_id(document: HwpxDocument, align: str) -> str | None:
    if not document.headers:
        return None
    return document.headers[0].ensure_paragraph_alignment(align.upper())


def _char_style(
    document: HwpxDocument,
    *,
    bold: bool = False,
    size: float = 10.5,
    font: str = "함초롬바탕",
    color: str = "#000000",
) -> str:
    return document.ensure_run_style(
        bold=bold,
        size=size,
        font=font,
        color=color,
    )


def _add_paragraph(
    document: HwpxDocument,
    text: str,
    *,
    bold: bool = False,
    size: float = 10.5,
    font: str = "함초롬바탕",
    align: str | None = None,
) -> None:
    para_id = _paragraph_alignment_id(document, align) if align else None
    document.add_paragraph(
        text,
        para_pr_id_ref=para_id,
        char_pr_id_ref=_char_style(
            document,
            bold=bold,
            size=size,
            font=font,
        ),
    )


def _date_to_korean(value: str) -> str:
    if not value:
        return ""
    parts = value.split("-")
    if len(parts) == 3 and all(part.isdigit() for part in parts):
        return f"{int(parts[0])}년 {int(parts[1])}월 {int(parts[2])}일"
    return value


def _render_notice_hwpx(item: GeneratedDocument) -> bytes:
    document = HwpxDocument.new()

    if item.notice_number:
        _add_paragraph(document, item.notice_number, size=9.5)

    _add_paragraph(
        document,
        item.title,
        bold=True,
        size=18,
        font="함초롬돋움",
        align="CENTER",
    )
    _add_paragraph(document, "", size=6)

    if item.alert_text:
        alert_border = document.ensure_border_fill(
            border_color="#666666",
            border_width="0.12 mm",
            fill_color="#F7F7F7",
        )
        alert = document.add_table(
            rows=1,
            cols=1,
            border_fill_id_ref=alert_border,
        )
        alert.set_cell_text(0, 0, item.alert_text)
        alert.set_cell_shading(0, 0, "F7F7F7")
        _add_paragraph(document, "", size=5)

    body_lines = item.content.splitlines()
    in_summary = False
    summary_inserted = False
    sub_index = 0
    korean = ("가", "나", "다", "라", "마", "바", "사", "아", "자", "차", "카", "타", "파", "하")

    for raw in body_lines:
        line = raw.strip()
        if not line:
            continue
        if line.startswith("# "):
            continue
        if line.startswith("- 공고번호:") or line.startswith("- 발주기관:"):
            continue

        if line.startswith("## "):
            heading = line[3:].strip()
            in_summary = bool(item.summary_rows and heading.startswith("1."))
            sub_index = 0
            _add_paragraph(
                document,
                heading,
                bold=True,
                size=12,
                font="함초롬돋움",
            )
            if in_summary and not summary_inserted:
                border = document.ensure_border_fill(
                    border_color="#666666",
                    border_width="0.12 mm",
                )
                table = document.add_table(
                    rows=len(item.summary_rows),
                    cols=2,
                    border_fill_id_ref=border,
                )
                table.set_column_widths([1.3, 4.7])
                for row_index, (label, value) in enumerate(item.summary_rows):
                    table.set_cell_text(row_index, 0, label)
                    table.set_cell_text(row_index, 1, value)
                    table.set_cell_shading(row_index, 0, "F2F2F2")
                summary_inserted = True
            continue

        if in_summary:
            continue

        if line.startswith("- "):
            prefix = korean[sub_index] + "." if sub_index < len(korean) else "•"
            sub_index += 1
            _add_paragraph(document, f"{prefix} {line[2:].strip()}", size=10.5)
            continue

        if line.startswith("※"):
            _add_paragraph(document, line, bold=True, size=9.5)
            continue

        _add_paragraph(document, line, size=10.5)

    if item.issue_date or item.signatory:
        _add_paragraph(document, "", size=8)
        _add_paragraph(document, "위와 같이 공고합니다.", size=11, align="CENTER")
        if item.issue_date:
            _add_paragraph(
                document,
                _date_to_korean(item.issue_date),
                size=11,
                align="CENTER",
            )
        if item.signatory:
            _add_paragraph(
                document,
                item.signatory,
                bold=True,
                size=16,
                font="함초롬돋움",
                align="CENTER",
            )

    document.set_footer_content(
        [
            {
                "align": "CENTER",
                "children": [
                    {
                        "type": "page_number",
                        "format": "page",
                        "position": "BOTTOM_CENTER",
                    }
                ],
            }
        ],
        section_index=0,
    )
    return document.to_bytes()


def document_to_hwpx_bytes(item: GeneratedDocument) -> bytes:
    if item.layout == "public_notice":
        return _render_notice_hwpx(item)

    document = HwpxDocument.new()
    for raw in item.content.splitlines():
        line = raw.rstrip()
        if not line:
            document.add_paragraph("")
        elif line.startswith("# "):
            document.add_heading(line[2:], level=1)
        elif line.startswith("## "):
            document.add_heading(line[3:], level=2)
        elif line.startswith("- "):
            document.add_paragraph("• " + line[2:])
        else:
            document.add_paragraph(line)
    return document.to_bytes()


def build_hwpx_zip(documents: list[GeneratedDocument]) -> bytes:
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        for index, item in enumerate(documents, start=1):
            archive.writestr(
                f"{index:02d}_{_safe_filename(item.title)}.hwpx",
                document_to_hwpx_bytes(item),
            )
    return stream.getvalue()


def _docx_font(document: Document) -> None:
    style = document.styles["Normal"]
    style.font.name = "Malgun Gothic"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    style.font.size = Pt(10.5)


def document_to_docx_bytes(item: GeneratedDocument) -> bytes:
    document = Document()
    _docx_font(document)
    for raw in item.content.splitlines():
        line = raw.rstrip()
        if not line:
            document.add_paragraph()
        elif line.startswith("# "):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:])
            run.bold = True
            run.font.size = Pt(16)
        elif line.startswith("## "):
            p = document.add_paragraph()
            run = p.add_run(line[3:])
            run.bold = True
            run.font.size = Pt(12)
        elif line.startswith("- "):
            document.add_paragraph(line[2:], style="List Bullet")
        else:
            document.add_paragraph(line)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def build_docx_zip(documents: list[GeneratedDocument]) -> bytes:
    stream = BytesIO()
    with ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive:
        for index, item in enumerate(documents, start=1):
            archive.writestr(
                f"{index:02d}_{_safe_filename(item.title)}.docx",
                document_to_docx_bytes(item),
            )
    return stream.getvalue()
